"""
PRUEBA 21 - Generación correcta de reportes (vista y documento PDF/Excel).

Comprueba que los datos aparezcan en las columnas indicadas, que los totales
sean correctos y que los documentos descargables sean archivos válidos.

Ejecutar:  pytest tests/test_reportes_documentos.py -v

Autor: 2024-06-10 - Vitalia S.A. - Equipo de Desarrollo - Guadarrama Hernandez Jose Luis

"""
from io import BytesIO

from openpyxl import load_workbook

from datetime import date, datetime, time
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.core.dependencies import get_current_user
from app.database.connection import get_db

client = TestClient(app)

ADMIN, DOCTOR, PACIENTE, RECEPCION, COORDINADOR = 1, 2, 3, 4, 5
NOMBRE_ROL = {1: "Administrador", 2: "Doctor", 3: "Paciente", 4: "Recepcionista", 5: "Coordinador"}


class _Consulta:
    """Simula la cadena query().filter().order_by()... de SQLAlchemy."""

    def __init__(self, valor):
        self.valor = valor

    def __getattr__(self, nombre):
        if nombre in ("filter", "order_by", "options", "join", "limit", "offset"):
            return lambda *a, **k: self
        raise AttributeError(nombre)

    def first(self):
        if isinstance(self.valor, list):
            return self.valor[0] if self.valor else None
        return self.valor

    def all(self):
        if self.valor is None:
            return []
        return self.valor if isinstance(self.valor, list) else [self.valor]

    def delete(self):
        return 0

    def count(self):
        return len(self.all())


class FakeDB:
    """Base de datos falsa: no necesita MySQL. `resultados` mapea Modelo -> valor."""

    def __init__(self, resultados=None):
        self.resultados = resultados or {}
        self.agregados = []
        self.commits = 0
        self.rollbacks = 0

    def _clave(self, modelo):
        return getattr(modelo, "class_", modelo)

    def query(self, *modelos):
        return _Consulta(self.resultados.get(self._clave(modelos[0])))

    def get(self, modelo, _id):
        return self.resultados.get(modelo)

    def add(self, obj):
        if hasattr(obj, "id_PK") and getattr(obj, "id_PK", None) is None:
            obj.id_PK = len(self.agregados) + 1
        self.agregados.append(obj)

    def flush(self):
        pass

    def commit(self):
        self.commits += 1

    def rollback(self):
        self.rollbacks += 1

    def refresh(self, obj):
        pass

    def delete(self, obj):
        pass


def usuario(rol_id, id_usuario=1, nombre="Usuario"):
    return SimpleNamespace(
        id_PK=id_usuario,
        nombre=nombre,
        correo_electronico=f"{nombre.lower()}@vitalia.com",
        rol_id_FK=rol_id,
        estado="activo",
    )


def entrar_como(rol_id, resultados=None, id_usuario=1):
    """Autentica como `rol_id` y devuelve la FakeDB inyectada en la API."""
    from app.models.role import Role

    res = {Role: SimpleNamespace(id_PK=rol_id, nombre=NOMBRE_ROL[rol_id], esta_activo=True)}
    res.update(resultados or {})
    db = FakeDB(res)
    app.dependency_overrides[get_current_user] = lambda: usuario(rol_id, id_usuario)
    app.dependency_overrides[get_db] = lambda: db
    return db


@pytest.fixture(autouse=True)
def limpiar_overrides():
    yield
    app.dependency_overrides.clear()

from app.models.cita import Cita
from app.models.paciente import Paciente
from app.services import reporte_export
from app.services.reporte_service import ReporteService


def _ejecucion(filas):
    m = MagicMock()
    m.all.return_value = filas
    return m


# ------------------------------------------------ CÁLCULOS (datos estadísticos)
def test_citas_por_doctor_totales_correctos():
    db = MagicMock()
    # 1ª consulta: nombres de doctores; 2ª: conteo por doctor/estado
    db.execute.side_effect = [
        _ejecucion([(1, "Dra. Ana"), (2, "Dr. Beto")]),
        _ejecucion([(1, "programada", 3), (1, "cancelada", 1), (2, "completada", 2)]),
    ]
    r = ReporteService(db).citas_por_doctor(None, None)
    assert r["total_citas"] == 6
    ana = next(x for x in r["detalle"] if x["doctor_id"] == 1)
    assert ana["doctor"] == "Dra. Ana"
    assert ana["total"] == 4
    assert ana["por_estado"] == {"programada": 3, "cancelada": 1}
    assert r["detalle"][0]["doctor_id"] == 1   # ordenado de mayor a menor


def test_tiempos_de_espera_promedio_y_maximo():
    db = MagicMock()
    base = datetime(2030, 1, 1, 9, 0)
    db.execute.side_effect = [
        _ejecucion([(1, "Dra. Ana")]),
        _ejecucion([
            (1, base, datetime(2030, 1, 1, 9, 10)),   # 10 min
            (1, base, datetime(2030, 1, 1, 9, 20)),   # 20 min
        ]),
    ]
    r = ReporteService(db).tiempos_espera(None, None)
    assert r["total_atenciones"] == 2
    assert r["promedio_minutos"] == 15
    assert r["detalle"][0]["maximo_minutos"] == 20


def test_ocupacion_de_salas_cuenta_horas_y_citas_de_la_sala_asignada():
    sala = SimpleNamespace(id_PK=1, nombre="Sala 1", tipo="consulta", piso="1",
                           estado="disponible", capacidad=4)
    db = MagicMock()
    db.scalars.return_value.all.return_value = [sala]
    db.execute.side_effect = [
        _ejecucion([]),                                   # citas por ubicación de calendario
        _ejecucion([(1, time(9, 0), time(11, 0)),         # citas con sala_id_FK
                    (1, time(12, 0), time(13, 0))]),
    ]
    r = ReporteService(db).ocupacion_salas(None, None)
    fila = r["detalle"][0]
    assert fila["total_citas"] == 2
    assert fila["horas_ocupadas"] == 3
    assert r["por_estado"] == {"disponible": 1}


# ------------------------------------------------------ DOCUMENTOS PDF / EXCEL
DATOS = {
    "ocupacion": {"desde": None, "hasta": None, "total_salas": 1, "por_estado": {"disponible": 1},
                  "detalle": [{"sala_id": 1, "nombre": "Sala 1", "tipo": "consulta", "piso": "1",
                               "estado": "disponible", "capacidad": 4, "total_citas": 2, "horas_ocupadas": 3.0}]},
    "citas": {"desde": "2030-01-01", "hasta": "2030-01-31", "total_citas": 4,
              "detalle": [{"doctor_id": 1, "doctor": "Dra. Ana", "total": 4,
                           "por_estado": {"programada": 3, "cancelada": 1}}]},
    "espera": {"desde": None, "hasta": None, "total_atenciones": 2, "promedio_minutos": 15,
               "detalle": [{"doctor_id": 1, "doctor": "Dra. Ana", "atenciones": 2,
                            "promedio_minutos": 15, "maximo_minutos": 20}]},
    "inventario": {"umbral": 10, "total": 1,
                   "detalle": [{"medicamento_id": 1, "nombre": "Paracetamol", "presentacion": "500 mg",
                                "categoria": "Analgésico", "stock_total": 4, "lotes": 1,
                                "proximo_vencimiento": "2031-01-01"}]},
    "recetas": {"desde": None, "hasta": None, "total": 3, "por_estado": {"emitida": 3},
                "detalle": [{"doctor_id": 1, "doctor": "Dra. Ana", "total": 3}]},
}


@pytest.mark.parametrize("tipo", list(DATOS))
def test_pdf_es_valido_para_cada_reporte(tipo):
    pdf = reporte_export.a_pdf(reporte_export.preparar(tipo, DATOS[tipo]))
    assert pdf.startswith(b"%PDF")
    assert len(pdf) > 1000


@pytest.mark.parametrize("tipo", list(DATOS))
def test_excel_tiene_datos_en_las_columnas_indicadas(tipo):
    rep = reporte_export.preparar(tipo, DATOS[tipo])
    libro = load_workbook(BytesIO(reporte_export.a_excel(rep)))
    hoja = libro.active
    encabezados = [c.value for c in next(
        f for f in hoja.iter_rows() if f[0].value == rep["columnas"][0])]
    assert encabezados[: len(rep["columnas"])] == rep["columnas"]
    valores = {c.value for f in hoja.iter_rows() for c in f}
    for fila in rep["filas"]:
        for dato in fila:
            if dato != "":
                assert dato in valores
    assert len(hoja._charts) == 1          # incluye gráfico


def test_excel_citas_columnas_por_estado():
    rep = reporte_export.preparar("citas", DATOS["citas"])
    assert rep["columnas"] == ["Doctor", "Total", "Cancelada", "Programada"]
    assert rep["filas"] == [["Dra. Ana", 4, 1, 3]]


def test_reporte_sin_datos_genera_documentos_sin_error():
    vacio = {"desde": None, "hasta": None, "total_citas": 0, "detalle": []}
    rep = reporte_export.preparar("citas", vacio)
    assert reporte_export.a_pdf(rep).startswith(b"%PDF")
    assert reporte_export.a_excel(rep)[:2] == b"PK"   # xlsx es un zip


# --------------------------------------------------------------- ENDPOINTS
def test_endpoint_exportar_pdf_y_excel():
    entrar_como(ADMIN)
    with patch("app.routers.reportes.ReporteService") as srv:
        srv.return_value.citas_por_doctor.return_value = DATOS["citas"]
        pdf = client.get("/api/reportes/citas/exportar?formato=pdf")
        xls = client.get("/api/reportes/citas/exportar?formato=excel")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert pdf.headers["content-type"] == "application/pdf"
    assert "reporte_citas.pdf" in pdf.headers["content-disposition"]
    assert xls.status_code == 200 and xls.content[:2] == b"PK"
    assert "reporte_citas.xlsx" in xls.headers["content-disposition"]


def test_endpoint_exportar_formato_o_tipo_invalido():
    entrar_como(ADMIN)
    assert client.get("/api/reportes/citas/exportar?formato=word").status_code == 422
    assert client.get("/api/reportes/inexistente/exportar?formato=pdf").status_code == 404


def test_endpoint_reporte_en_pantalla_devuelve_los_mismos_datos():
    entrar_como(ADMIN)
    with patch("app.routers.reportes.ReporteService") as srv:
        srv.return_value.citas_por_doctor.return_value = DATOS["citas"]
        r = client.get("/api/reportes/citas-por-doctor")
    assert r.status_code == 200
    assert r.json()["total_citas"] == 4


@pytest.mark.parametrize("rol", [DOCTOR, PACIENTE, RECEPCION])
def test_reportes_restringidos_a_admin_y_coordinador(rol):
    entrar_como(rol)
    assert client.get("/api/reportes/citas-por-doctor").status_code == 403
    assert client.get("/api/reportes/citas/exportar?formato=pdf").status_code == 403


def test_reportes_sin_token_401():
    assert client.get("/api/reportes/citas-por-doctor").status_code == 401


# ------------------------------------------------- DATOS DEL PACIENTE
def test_paciente_ve_sus_citas_con_sala_y_estado():
    pac = SimpleNamespace(id_PK=7)
    cita = Cita(id_PK=1, paciente_id_FK=7, doctor_id_FK=1, fecha=date(2020, 1, 1),
                hora_inicio=time(9), hora_fin=time(9, 30), estado="no_asistio",
                motivo="Chequeo", requiere_equipo=False, creado_por=1)
    entrar_como(PACIENTE, {Paciente: pac, Cita: [cita]}, id_usuario=3)
    r = client.get("/api/citas/mis-citas")
    assert r.status_code == 200, r.text
    fila = r.json()[0]
    assert fila["estado"] == "no_asistio"
    assert fila["no_realizada"] is True
    assert fila["motivo"] == "Chequeo"


def test_solo_paciente_puede_ver_mis_citas():
    entrar_como(RECEPCION)
    assert client.get("/api/citas/mis-citas").status_code == 403

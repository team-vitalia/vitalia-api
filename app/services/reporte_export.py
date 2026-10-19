"""Exportación de reportes a PDF (con gráfico) y Excel (con gráfico nativo)."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

VERDE = "247F76"

TIPOS = ("ocupacion", "citas", "espera", "inventario", "recetas")


def _doctor(x: dict) -> str:
    if x.get("doctor"):
        return x["doctor"]
    return f"Doctor #{x['doctor_id']}" if x.get("doctor_id") is not None else "Sin doctor"


def _periodo(data: dict) -> str:
    d, h = data.get("desde"), data.get("hasta")
    if not d and not h:
        return "Periodo: todo el historial"
    return f"Periodo: {d or 'inicio'} a {h or 'hoy'}"


def preparar(tipo: str, data: dict) -> dict:
    """Convierte la respuesta de un reporte en tabla + serie para el gráfico."""
    if tipo == "ocupacion":
        return {
            "titulo": "Ocupación de salas",
            "subtitulo": _periodo(data),
            "resumen": [("Salas registradas", data["total_salas"])]
            + [(f"Estado «{k}»", v) for k, v in data["por_estado"].items()],
            "columnas": ["Sala", "Tipo", "Piso", "Estado", "Capacidad", "Citas", "Horas ocupadas"],
            "filas": [
                [x["nombre"], x["tipo"] or "", x["piso"] or "", x["estado"] or "",
                 x["capacidad"] or "", x["total_citas"], x["horas_ocupadas"]]
                for x in data["detalle"]
            ],
            "grafico": ("Horas ocupadas", [x["nombre"] for x in data["detalle"]],
                        [x["horas_ocupadas"] for x in data["detalle"]]),
            "col_grafico": 6,
        }
    if tipo == "citas":
        estados = sorted({k for x in data["detalle"] for k in x["por_estado"]})
        return {
            "titulo": "Citas por doctor",
            "subtitulo": _periodo(data),
            "resumen": [("Total de citas", data["total_citas"])],
            "columnas": ["Doctor", "Total"] + [e.replace("_", " ").capitalize() for e in estados],
            "filas": [[_doctor(x), x["total"]] + [x["por_estado"].get(e, 0) for e in estados]
                      for x in data["detalle"]],
            "grafico": ("Citas", [_doctor(x) for x in data["detalle"]], [x["total"] for x in data["detalle"]]),
            "col_grafico": 1,
        }
    if tipo == "espera":
        return {
            "titulo": "Tiempos de espera",
            "subtitulo": _periodo(data),
            "resumen": [("Espera promedio (min)", data["promedio_minutos"]),
                        ("Pacientes atendidos", data["total_atenciones"])],
            "columnas": ["Doctor", "Atenciones", "Promedio (min)", "Máximo (min)"],
            "filas": [[_doctor(x), x["atenciones"], x["promedio_minutos"], x["maximo_minutos"]]
                      for x in data["detalle"]],
            "grafico": ("Promedio (min)", [_doctor(x) for x in data["detalle"]],
                        [x["promedio_minutos"] for x in data["detalle"]]),
            "col_grafico": 2,
        }
    if tipo == "inventario":
        return {
            "titulo": "Inventario bajo",
            "subtitulo": f"Medicamentos con {data['umbral']} unidades o menos",
            "resumen": [("Medicamentos", data["total"])],
            "columnas": ["Medicamento", "Presentación", "Categoría", "Lotes", "Próx. vencimiento", "Stock"],
            "filas": [[x["nombre"], x["presentacion"] or "", x["categoria"] or "", x["lotes"],
                       x["proximo_vencimiento"] or "", x["stock_total"]] for x in data["detalle"]],
            "grafico": ("Stock", [x["nombre"] for x in data["detalle"]],
                        [x["stock_total"] for x in data["detalle"]]),
            "col_grafico": 5,
        }
    if tipo == "recetas":
        return {
            "titulo": "Recetas emitidas",
            "subtitulo": _periodo(data),
            "resumen": [("Recetas emitidas", data["total"])]
            + [(f"Estado «{k}»", v) for k, v in data["por_estado"].items()],
            "columnas": ["Doctor", "Recetas"],
            "filas": [[_doctor(x), x["total"]] for x in data["detalle"]],
            "grafico": ("Recetas", [_doctor(x) for x in data["detalle"]], [x["total"] for x in data["detalle"]]),
            "col_grafico": 1,
        }
    raise ValueError(tipo)


def a_excel(rep: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte"
    ws["A1"] = f"VITALIA · {rep['titulo']}"
    ws["A1"].font = Font(bold=True, size=14, color=VERDE)
    ws["A2"] = rep["subtitulo"]
    fila = 3
    for k, v in rep["resumen"]:
        ws.cell(row=fila, column=1, value=k).font = Font(bold=True)
        ws.cell(row=fila, column=2, value=v)
        fila += 1
    fila += 1
    cab = fila
    for c, nombre in enumerate(rep["columnas"], 1):
        cel = ws.cell(row=cab, column=c, value=nombre)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor=VERDE)
        cel.alignment = Alignment(horizontal="center")
    for r in rep["filas"]:
        fila += 1
        for c, v in enumerate(r, 1):
            ws.cell(row=fila, column=c, value=v)
    for c in range(1, len(rep["columnas"]) + 1):
        ancho = max(len(str(ws.cell(row=r, column=c).value or "")) for r in range(cab, fila + 1))
        ws.column_dimensions[get_column_letter(c)].width = min(max(ancho + 3, 12), 40)

    if rep["filas"]:
        ch = BarChart()
        ch.type = "col"
        ch.title = rep["titulo"]
        ch.y_axis.title = rep["grafico"][0]
        col = rep["col_grafico"] + 1
        ch.add_data(Reference(ws, min_col=col, min_row=cab, max_row=fila), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=cab + 1, max_row=fila))
        ch.legend = None
        ch.height, ch.width = 9, 20
        ws.add_chart(ch, f"A{fila + 3}")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def a_pdf(rep: dict) -> bytes:
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=1.5 * cm, rightMargin=1.5 * cm,
                            topMargin=1.5 * cm, bottomMargin=1.5 * cm, title=rep["titulo"])
    est = getSampleStyleSheet()
    verde = colors.HexColor(f"#{VERDE}")
    el = [Paragraph(f"VITALIA · {rep['titulo']}", est["Title"]), Paragraph(rep["subtitulo"], est["Normal"]),
          Spacer(1, 8)]
    for k, v in rep["resumen"]:
        el.append(Paragraph(f"<b>{k}:</b> {v}", est["Normal"]))
    el.append(Spacer(1, 10))

    if rep["filas"]:
        nombre, etiquetas, valores = rep["grafico"]
        etiquetas, valores = etiquetas[:12], valores[:12]
        d = Drawing(700, 220)
        bc = VerticalBarChart()
        bc.x, bc.y, bc.width, bc.height = 50, 40, 620, 150
        bc.data = [[float(v) for v in valores]]
        bc.categoryAxis.categoryNames = [(e or "")[:14] for e in etiquetas]
        bc.categoryAxis.labels.angle = 25
        bc.categoryAxis.labels.boxAnchor = "ne"
        bc.categoryAxis.labels.fontSize = 7
        bc.valueAxis.valueMin = 0
        bc.valueAxis.labels.fontSize = 7
        bc.bars[0].fillColor = verde
        d.add(bc)
        el += [Paragraph(f"<b>{nombre}</b>", est["Heading4"]), d, Spacer(1, 10)]

        tabla = Table([rep["columnas"]] + [[str(c) for c in f] for f in rep["filas"]], repeatRows=1)
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), verde), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F7F6")]),
        ]))
        el.append(tabla)
    else:
        el.append(Paragraph("Sin datos para el periodo seleccionado.", est["Normal"]))
    doc.build(el)
    return buf.getvalue()
